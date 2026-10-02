"""LocalTextParser: txt / eml / csv / xlsx / pdf-text-layer, in-process, no I/O beyond the bytes.

Hard rules: no network, no link fetching, nothing is executed or evaluated (xlsx formulas are
never computed; cached values only), size/page/row limits, all output text goes through
``inert_text``. PDFs are only read when ``pypdf`` is installed; otherwise the result is flagged
``unsupported:pdf_needs_sandbox`` (stack-map S4).
"""

from __future__ import annotations

import csv
import hashlib
import io
import zipfile
from email import policy
from email.message import EmailMessage, Message
from email.parser import BytesParser
from typing import Any

from components.rfq.quotes.inert import LINK_REMOVED, inert_text, reveal_text

from .models import DocumentParser, ParsedDocument, ParserLimits, SourceSpan, Table

__all__ = ["LocalTextParser", "DocumentParser"]

_ACTIVE_PDF_MARKERS = (b"/JavaScript", b"/JS", b"/Launch", b"/EmbeddedFile", b"/OpenAction")


class _Builder:
    """Accumulates inert segments into one text with offsets, honouring max_chars."""

    def __init__(self, limits: ParserLimits) -> None:
        self.limits = limits
        self.parts: list[str] = []
        self.spans: list[SourceSpan] = []
        self.flags: list[str] = []
        self.size = 0

    def add(self, raw: str, part: str, page: int | None = None) -> str:
        clean = inert_text(raw)
        self._note(raw, clean)
        if not clean:
            return ""
        sep = 2 if self.parts else 0
        room = self.limits.max_chars - self.size - sep
        if room <= 0:
            self.flag("truncated:max_chars")
            return ""
        if len(clean) > room:
            clean = clean[:room].rstrip()
            self.flag("truncated:max_chars")
        start = self.size + sep
        self.parts.append(clean)
        self.size = start + len(clean)
        self.spans.append(SourceSpan(start, self.size, part, page))
        return clean

    def _note(self, raw: str, clean: str) -> None:
        if LINK_REMOVED in clean and LINK_REMOVED not in raw:
            self.flag("neutralised:links_removed")
        if inert_text(reveal_text(raw)) != clean:
            self.flag("neutralised:hidden_or_active_content")

    def flag(self, f: str) -> None:
        if f not in self.flags:
            self.flags.append(f)

    @property
    def text(self) -> str:
        return "\n\n".join(self.parts)


def _decode(content: bytes) -> str:
    if content.startswith(b"\xef\xbb\xbf"):
        content = content[3:]
    return content.decode("utf-8", errors="replace")


def _cell(v: object, limits: ParserLimits) -> str:
    s = "" if v is None else str(v)
    return inert_text(s[: limits.max_cell_chars])


def _table_text(rows: list[tuple[str, ...]]) -> str:
    return "\n".join(" | ".join(r) for r in rows)


class LocalTextParser:
    def __init__(self, limits: ParserLimits | None = None) -> None:
        self.limits = limits or ParserLimits()

    # ------------------------------------------------------------------ dispatch

    def parse(
        self, content: bytes, *, filename: str, content_type: str | None = None
    ) -> ParsedDocument:
        name = inert_text(filename.replace("\\", "/").rsplit("/", 1)[-1])[:255]
        digest = hashlib.sha256(content).hexdigest()

        def rejected(flag: str) -> ParsedDocument:
            return ParsedDocument("", filename=name, sha256=digest, flags=(flag,))

        if not isinstance(content, bytes | bytearray):
            raise TypeError("content must be bytes")
        if len(content) > self.limits.max_bytes:
            return rejected("rejected:too_large")
        kind = self._kind(bytes(content), name, content_type)
        try:
            if kind == "pdf":
                return self._pdf(bytes(content), name, digest)
            if kind == "xlsx":
                return self._xlsx(bytes(content), name, digest)
            if kind == "eml":
                return self._eml(bytes(content), name, digest)
            if kind == "csv":
                return self._csv(bytes(content), name, digest)
            if kind == "txt":
                return self._txt(bytes(content), name, digest)
        except Exception as exc:  # noqa: BLE001 - hostile input must never crash the caller
            return rejected(f"error:parse_failed:{type(exc).__name__}")
        return rejected("unsupported:file_type")

    @staticmethod
    def _kind(content: bytes, name: str, content_type: str | None) -> str:
        ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
        if content[:5] == b"%PDF-" or ext == "pdf":
            return "pdf"
        if content[:2] == b"PK" and ext in ("xlsx", "xlsm", "zip", ""):
            return "xlsx" if ext != "zip" else "other"
        if ext in ("xlsx", "xlsm"):
            return "xlsx"
        if ext == "eml" or (content_type or "").lower() == "message/rfc822":
            return "eml"
        if ext in ("csv", "tsv") or (content_type or "").lower() == "text/csv":
            return "csv"
        if ext in ("txt", "text", "") or (content_type or "").lower() == "text/plain":
            return "txt"
        return "other"

    # ------------------------------------------------------------------ txt / csv

    def _txt(self, content: bytes, name: str, digest: str) -> ParsedDocument:
        b = _Builder(self.limits)
        b.add(_decode(content), "body")
        return ParsedDocument(b.text, (b.text,), (), tuple(b.spans), tuple(b.flags), name, digest)

    def _csv(self, content: bytes, name: str, digest: str) -> ParsedDocument:
        text = _decode(content)
        delim = "\t" if name.lower().endswith(".tsv") else ","
        flags: list[str] = []
        rows: list[tuple[str, ...]] = []
        csv.field_size_limit(self.limits.max_cell_chars * 4)
        try:
            for i, rec in enumerate(csv.reader(io.StringIO(text, newline=""), delimiter=delim)):
                if i >= self.limits.max_rows:
                    flags.append("truncated:max_rows")
                    break
                if len(rec) > self.limits.max_cols:
                    rec = rec[: self.limits.max_cols]
                    if "truncated:max_cols" not in flags:
                        flags.append("truncated:max_cols")
                rows.append(tuple(_cell(c, self.limits) for c in rec))
        except csv.Error:
            flags.append("error:malformed_csv")
        return self._table_doc([("table", rows)], name, digest, flags)

    def _table_doc(
        self, sheets: list[tuple[str, list[tuple[str, ...]]]], name: str, digest: str,
        flags: list[str],
    ) -> ParsedDocument:
        b = _Builder(self.limits)
        b.flags.extend(flags)
        tables: list[Table] = []
        pages: list[str] = []
        for idx, (title, rows) in enumerate(sheets, start=1):
            tables.append(tuple(rows))
            added = b.add(_table_text(rows), f"sheet:{title}" if len(sheets) > 1 or
                          title != "table" else "table", idx)
            pages.append(added)
        return ParsedDocument(
            b.text, tuple(pages), tuple(tables), tuple(b.spans), tuple(b.flags), name, digest
        )

    # ------------------------------------------------------------------ eml

    def _eml(self, content: bytes, name: str, digest: str) -> ParsedDocument:
        msg = BytesParser(policy=policy.default).parsebytes(content)
        b = _Builder(self.limits)
        if msg.defects:
            b.flag("warn:malformed_eml")
        for hdr in ("From", "To", "Subject", "Date"):
            try:
                val = str(msg.get(hdr, "") or "")
            except Exception:  # noqa: BLE001 - malformed header objects
                val, _ = "", b.flag("warn:malformed_eml")
            if val:
                b.add(f"{hdr}: {val[:1000]}", f"header:{hdr.lower()}")
        plain: list[str] = []
        html: list[str] = []
        attachments: list[str] = []
        for n, part in enumerate(msg.walk()):
            if n >= self.limits.max_parts:
                b.flag("truncated:max_parts")
                break
            if part.is_multipart():
                continue
            self._eml_part(part, plain, html, attachments, b)
        body = plain or html
        if not plain and html:
            b.flag("body:html_only")
        for i, chunk in enumerate(body):
            b.add(chunk[: self.limits.max_chars], "body", i + 1)
        if attachments:
            b.add("Attachments:\n" + "\n".join(attachments), "attachments")
        if not body:
            b.flag("warn:empty_body")
        return ParsedDocument(
            b.text, (), (), tuple(b.spans), tuple(b.flags), name, digest,
            tuple(inert_text(a) for a in attachments),
        )

    def _eml_part(
        self, part: Message, plain: list[str], html: list[str], attachments: list[str],
        b: _Builder,
    ) -> None:
        ctype = part.get_content_type()
        fname = part.get_filename()
        if fname or part.get_content_disposition() == "attachment":
            try:
                size = len(part.get_payload(decode=True) or b"")
            except Exception:  # noqa: BLE001
                size = 0
            label = inert_text((fname or "unnamed").replace("\\", "/").rsplit("/", 1)[-1])[:120]
            attachments.append(f"{label or 'unnamed'} ({ctype}, {size} bytes)")
            b.flag("attachments:listed_not_parsed")
            return
        if ctype not in ("text/plain", "text/html"):
            return
        try:
            assert isinstance(part, EmailMessage)
            body = part.get_content()
        except Exception:  # noqa: BLE001 - unknown charset etc.
            payload: Any = part.get_payload(decode=True) or b""
            body = payload.decode("utf-8", errors="replace")
            b.flag("warn:malformed_eml")
        (plain if ctype == "text/plain" else html).append(str(body))

    # ------------------------------------------------------------------ xlsx

    def _xlsx(self, content: bytes, name: str, digest: str) -> ParsedDocument:
        try:
            import openpyxl
        except ImportError:
            return ParsedDocument("", filename=name, sha256=digest,
                                  flags=("unsupported:xlsx_needs_openpyxl",))
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as zf:
                infos = zf.infolist()
                if len(infos) > 2000 or sum(i.file_size for i in infos) > (
                    self.limits.max_unzipped_bytes
                ):
                    return ParsedDocument("", filename=name, sha256=digest,
                                          flags=("rejected:zip_bomb",))
                for info in infos:
                    if info.filename.endswith(".xml") and info.file_size < 5_000_000:
                        head = zf.read(info)
                        if b"<!ENTITY" in head or b"<!DOCTYPE" in head:
                            return ParsedDocument("", filename=name, sha256=digest,
                                                  flags=("rejected:xml_entities",))
        except zipfile.BadZipFile:
            return ParsedDocument("", filename=name, sha256=digest, flags=("error:bad_zip",))
        wb = openpyxl.load_workbook(
            io.BytesIO(content), read_only=True, data_only=True, keep_links=False
        )
        flags: list[str] = []
        sheets: list[tuple[str, list[tuple[str, ...]]]] = []
        try:
            for ws in wb.worksheets:
                if len(sheets) >= self.limits.max_pages:
                    flags.append("truncated:max_pages")
                    break
                if ws.sheet_state != "visible":
                    flags.append("hidden_sheet_skipped")
                    continue
                rows: list[tuple[str, ...]] = []
                for r in ws.iter_rows(values_only=True, max_col=self.limits.max_cols):
                    if len(rows) >= self.limits.max_rows:
                        flags.append("truncated:max_rows")
                        break
                    rows.append(tuple(_cell(c, self.limits) for c in r))
                sheets.append((inert_text(ws.title)[:60] or "sheet", rows))
        finally:
            wb.close()
        return self._table_doc(sheets, name, digest, flags)

    # ------------------------------------------------------------------ pdf

    def _pdf(self, content: bytes, name: str, digest: str) -> ParsedDocument:
        try:
            import pypdf
        except ImportError:
            return ParsedDocument("", filename=name, sha256=digest,
                                  flags=("unsupported:pdf_needs_sandbox",))
        flags = ["parsed_in_process"]
        if any(m in content for m in _ACTIVE_PDF_MARKERS):
            flags.append("neutralised:pdf_active_content_ignored")
        try:
            reader = pypdf.PdfReader(io.BytesIO(content), strict=False)
            if reader.is_encrypted:
                return ParsedDocument("", filename=name, sha256=digest,
                                      flags=("unsupported:pdf_encrypted",))
            total = len(reader.pages)
            if total > self.limits.max_pages:
                return ParsedDocument("", filename=name, sha256=digest,
                                      flags=("rejected:too_many_pages",))
            raw_pages = [(p.extract_text() or "") for p in reader.pages]
        except Exception as exc:  # noqa: BLE001
            return ParsedDocument("", filename=name, sha256=digest,
                                  flags=(f"error:pdf_unreadable:{type(exc).__name__}",))
        b = _Builder(self.limits)
        b.flags.extend(flags)
        pages: list[str] = []
        for i, raw in enumerate(raw_pages, start=1):
            pages.append(b.add(raw, f"page:{i}", i))
        if not any(pages):
            b.flag("pdf_no_text_layer")
        return ParsedDocument(b.text, tuple(pages), (), tuple(b.spans), tuple(b.flags), name, digest)


